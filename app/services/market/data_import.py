"""
Admin bulk job dataset import (CSV / XLSX / JSON).

Expected columns/fields (all optional except title):
  title, company, location, remote_status, employment_type,
  salary_min, salary_max, salary_currency, salary_period,
  experience_years_min, experience_years_max, education_level,
  skills, preferred_skills

skills / preferred_skills are semicolon-separated skill names, e.g.
"Python;Docker;AWS". Each row becomes one Job with status="completed"
directly (no need to run the text-parsing pipeline — the data already
arrives structured), and its own raw_text is left as a compact summary
for reference. Unmatched skills are handled the same way as everywhere
else in the app: matched against the taxonomy, or created as a
user-suggested entry for admin triage — never silently dropped.
"""
import csv
import io
import json

from openpyxl import load_workbook

from app.extensions import db
from app.models import Job, JobSkill
from app.services.skills.normalizer import get_or_suggest_skill

REQUIRED_FIELD = "title"
VALID_REMOTE_STATUSES = {"remote", "hybrid", "onsite", "unknown"}
VALID_EMPLOYMENT_TYPES = {"full_time", "part_time", "contract", "internship", "temporary", "unknown"}
VALID_EDUCATION_LEVELS = {"high_school", "associate", "bachelor", "master", "phd", "bootcamp", "self_taught", "other"}


class ImportParseError(Exception):
    pass


def parse_csv(file_bytes: bytes) -> list[dict]:
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportParseError(f"Could not decode CSV as UTF-8: {exc}") from exc
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def parse_xlsx(file_bytes: bytes) -> list[dict]:
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportParseError(f"Could not read XLSX file: {exc}") from exc

    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        return []

    rows = []
    for raw_row in rows_iter:
        row = {headers[i]: raw_row[i] for i in range(min(len(headers), len(raw_row)))}
        if any(v not in (None, "") for v in row.values()):
            rows.append(row)
    return rows


def parse_json(file_bytes: bytes) -> list[dict]:
    try:
        data = json.loads(file_bytes.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ImportParseError(f"Could not parse JSON: {exc}") from exc

    if isinstance(data, dict):
        data = data.get("jobs", [])
    if not isinstance(data, list):
        raise ImportParseError("JSON must be a list of job objects (or {\"jobs\": [...]}).")
    return data


def _clean_str(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _clean_int(value):
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def _clean_enum(value, valid_set, default):
    text = _clean_str(value)
    if text and text.lower() in valid_set:
        return text.lower()
    return default


def _split_skills(value):
    if value is None:
        return []
    if isinstance(value, list):
        return [str(s).strip() for s in value if str(s).strip()]
    text = str(value).strip()
    if not text:
        return []
    separator = ";" if ";" in text else ","
    return [s.strip() for s in text.split(separator) if s.strip()]


def _row_summary(cleaned: dict) -> str:
    parts = [cleaned.get("title") or "Untitled role"]
    if cleaned.get("company"):
        parts.append(f"at {cleaned['company']}")
    if cleaned.get("location"):
        parts.append(f"in {cleaned['location']}")
    return " ".join(parts) + " (imported dataset entry)"


def clean_row(row: dict):
    """Returns a cleaned dict ready for Job creation, or None if the row
    is unusable (missing the one required field)."""
    title = _clean_str(row.get("title"))
    if not title:
        return None

    cleaned = {
        "title": title,
        "company": _clean_str(row.get("company")),
        "location": _clean_str(row.get("location")),
        "remote_status": _clean_enum(row.get("remote_status"), VALID_REMOTE_STATUSES, "unknown"),
        "employment_type": _clean_enum(row.get("employment_type"), VALID_EMPLOYMENT_TYPES, "unknown"),
        "salary_min": _clean_int(row.get("salary_min")),
        "salary_max": _clean_int(row.get("salary_max")),
        "salary_currency": _clean_str(row.get("salary_currency")),
        "salary_period": _clean_str(row.get("salary_period")),
        "experience_years_min": _clean_int(row.get("experience_years_min")),
        "experience_years_max": _clean_int(row.get("experience_years_max")),
        "education_level": (
            str(row.get("education_level")).strip().lower()
            if row.get("education_level") and str(row.get("education_level")).strip().lower() in VALID_EDUCATION_LEVELS
            else None
        ),
        "skills": _split_skills(row.get("skills")),
        "preferred_skills": _split_skills(row.get("preferred_skills")),
    }

    if cleaned["salary_min"] is not None and cleaned["salary_min"] < 0:
        cleaned["salary_min"] = None
    if cleaned["salary_max"] is not None and cleaned["salary_max"] < 0:
        cleaned["salary_max"] = None
    if (
        cleaned["salary_min"] is not None
        and cleaned["salary_max"] is not None
        and cleaned["salary_min"] > cleaned["salary_max"]
    ):
        cleaned["salary_min"], cleaned["salary_max"] = cleaned["salary_max"], cleaned["salary_min"]

    return cleaned


def import_jobs(rows, admin_user_id: int) -> dict:
    """Clean, deduplicate, and bulk-insert Job rows from parsed dataset
    rows. Returns {"imported": n, "skipped": n, "duplicates": n, "errors": [...]}.
    """
    stats = {"imported": 0, "skipped": 0, "duplicates": 0, "errors": []}

    for index, raw_row in enumerate(rows, start=1):
        cleaned = clean_row(raw_row)
        if cleaned is None:
            stats["skipped"] += 1
            stats["errors"].append(f"Row {index}: missing required field 'title' — skipped.")
            continue

        existing = Job.query.filter_by(
            title=cleaned["title"], company=cleaned["company"], location=cleaned["location"], source="imported"
        ).first()
        if existing:
            stats["duplicates"] += 1
            continue

        job = Job(
            user_id=admin_user_id,
            title=cleaned["title"],
            company=cleaned["company"],
            location=cleaned["location"],
            remote_status=cleaned["remote_status"],
            employment_type=cleaned["employment_type"],
            salary_min=cleaned["salary_min"],
            salary_max=cleaned["salary_max"],
            salary_currency=cleaned["salary_currency"],
            salary_period=cleaned["salary_period"],
            experience_years_min=cleaned["experience_years_min"],
            experience_years_max=cleaned["experience_years_max"],
            education_level=cleaned["education_level"],
            source="imported",
            status="completed",
        )
        job.raw_text = _row_summary(cleaned)
        db.session.add(job)
        db.session.flush()

        seen_skill_ids = set()
        for skill_name in cleaned["skills"]:
            skill = get_or_suggest_skill(skill_name)
            if skill.id in seen_skill_ids:
                continue
            db.session.add(
                JobSkill(job_id=job.id, skill_id=skill.id, requirement_level="required", raw_text=skill_name)
            )
            seen_skill_ids.add(skill.id)

        for skill_name in cleaned["preferred_skills"]:
            skill = get_or_suggest_skill(skill_name)
            if skill.id in seen_skill_ids:
                continue
            db.session.add(
                JobSkill(job_id=job.id, skill_id=skill.id, requirement_level="preferred", raw_text=skill_name)
            )
            seen_skill_ids.add(skill.id)

        stats["imported"] += 1

    db.session.commit()
    return stats
